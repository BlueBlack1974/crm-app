from flask import Blueprint, jsonify, request, session, render_template, url_for, flash, redirect, send_file
from app.extensions import db
from app.models import Randevu, RandevuYetki, RandevuDefterAyar, Kullanici
from app.utils.decorators import login_required
from app.utils.logging import log_user_action
from datetime import datetime, timedelta
from collections import Counter, defaultdict
from io import StringIO, BytesIO
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

rapor_bp = Blueprint('rapor', __name__)

# Raporlar Sayfasi
@rapor_bp.route('/raporlar')
@login_required
def raporlar():
    # Modül izin kontrolü
    if not session.get('raporlar_modulu', False):
        flash('Bu sayfaya erişim yetkiniz yok', 'error')
        return redirect(url_for('main.dashboard'))
    # Varsayilan tarih araligi: son 30 gun
    end_str = request.args.get('bitis')
    start_str = request.args.get('baslangic')
    today = datetime.now().date()
    default_start = today - timedelta(days=30)
    try:
        start_date = datetime.strptime(start_str, '%Y-%m-%d').date() if start_str else default_start
    except ValueError:
        start_date = default_start
    try:
        end_date = datetime.strptime(end_str, '%Y-%m-%d').date() if end_str else today
    except ValueError:
        end_date = today

    # Kapasite ve defter parametrelerini al
    kapasite = request.args.get('kapasite', '')
    defter_id = request.args.get('defter_id', '')
    format_tip = request.args.get('format')  # Export format
    print(f"Form parametreleri - defter_id: '{defter_id}', kapasite: '{kapasite}', format: '{format_tip}'")
    
    # Defter listesini al
    defterler = RandevuDefterAyar.query.filter(
        RandevuDefterAyar.FirmaID == session['firma_id'],
        RandevuDefterAyar.Aktif == True
    ).all()
    print(f"Bulunan defter sayısı: {len(defterler)}")
    for defter in defterler:
        print(f"Defter: {defter.DefterAdi} (ID: {defter.AyarID})")
    
    # Eğer belirli bir defter seçilmişse ve kapasite boş/0 ise, defterden otomatik hesapla
    auto_capacity = None
    if defter_id and (not kapasite or kapasite == '0'):
        try:
            ayar = RandevuDefterAyar.query.filter_by(AyarID=defter_id, FirmaID=session['firma_id'], Aktif=True).first()
            if ayar:
                def parse_hhmm(s):
                    try:
                        h, m = (s or '09:00').split(':')
                        return int(h), int(m)
                    except:
                        return 9, 0  # Varsayılan 09:00
                
                sh, sm = parse_hhmm(ayar.BaslangicSaati or '09:00')
                eh, em = parse_hhmm(ayar.BitisSaati or '18:00')
                total_minutes = max(0, (eh * 60 + em) - (sh * 60 + sm))
                slot_min = max(1, ayar.SlotDakika or 30)  # En az 1 dakika
                auto_capacity = max(1, total_minutes // slot_min)
                print(f"Defter {ayar.DefterAdi} için otomatik kapasite hesaplandı: {auto_capacity} (Saat: {ayar.BaslangicSaati}-{ayar.BitisSaati}, Slot: {slot_min}dk)")
        except Exception as e:
            print(f"Kapasite hesaplama hatası: {e}")
            auto_capacity = None

    if not kapasite or kapasite == '0':
        kapasite = str(auto_capacity or 8)

    # Export işlemleri
    if format_tip in ['csv', 'excel', 'pdf']:
        # Randevu verilerini al
        start_dt = datetime(start_date.year, start_date.month, start_date.day, 0, 0)
        end_dt = datetime(end_date.year, end_date.month, end_date.day, 23, 59, 59)
        
        # Randevu sorgusu
        if session.get('is_admin', False):
            q = Randevu.query.filter(
                Randevu.FirmaID == session['firma_id'],
                Randevu.RandevuTarihi >= start_dt,
                Randevu.RandevuTarihi <= end_dt
            )
        else:
            q = db.session.query(Randevu).join(RandevuYetki).filter(
                RandevuYetki.KullaniciID == session['user_id'],
                RandevuYetki.GoruntulemeYetkisi == True,
                Randevu.FirmaID == session['firma_id'],
                Randevu.RandevuTarihi >= start_dt,
                Randevu.RandevuTarihi <= end_dt
            )
        
        # Defter filtresi
        if defter_id:
            q = q.filter(Randevu.DefterID == defter_id)
        
        randevular = q.order_by(Randevu.RandevuTarihi.desc()).all()
        
        # Çıktı alma logla
        try:
            log_user_action('VIEW', 'Rapor', detail=f"Randevu raporu {format_tip.upper()} indirildi")
        except Exception:
            pass
        
        if format_tip == 'csv':
            # CSV Export
            output = StringIO()
            writer = csv.writer(output, delimiter=';')
            writer.writerow(['Randevu ID', 'Tarih', 'Saat', 'Müşteri', 'Telefon', 'Email', 'Başlık', 'Durum', 'Süre', 'Defter'])
            for r in randevular:
                writer.writerow([
                    r.RandevuID,
                    r.RandevuTarihi.strftime('%d.%m.%Y'),
                    r.RandevuTarihi.strftime('%H:%M'),
                    f"{r.MusteriAdi} {r.MusteriSoyadi or ''}".strip(),
                    r.MusteriTelefon or '',
                    r.MusteriEmail or '',
                    r.RandevuBaslik or '',
                    r.Durum or '',
                    f"{r.RandevuSuresi or 60} dk",
                    r.defter.DefterAdi if r.defter else ''
                ])
            output.seek(0)
            filename = f"randevu_raporu_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
            data = output.getvalue().encode('utf-8-sig')
            return send_file(BytesIO(data), mimetype='text/csv; charset=utf-8', as_attachment=True, download_name=filename)
        
        elif format_tip == 'excel':
            # Excel Export
            wb = Workbook()
            ws = wb.active
            ws.title = "Randevu Raporu"
            
            # Başlık satırı
            headers = ['Randevu ID', 'Tarih', 'Saat', 'Müşteri', 'Telefon', 'Email', 'Başlık', 'Durum', 'Süre', 'Defter']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            
            # Veri satırları
            for row, r in enumerate(randevular, 2):
                ws.cell(row=row, column=1, value=r.RandevuID)
                ws.cell(row=row, column=2, value=r.RandevuTarihi.strftime('%d.%m.%Y'))
                ws.cell(row=row, column=3, value=r.RandevuTarihi.strftime('%H:%M'))
                ws.cell(row=row, column=4, value=f"{r.MusteriAdi} {r.MusteriSoyadi or ''}".strip())
                ws.cell(row=row, column=5, value=r.MusteriTelefon or '')
                ws.cell(row=row, column=6, value=r.MusteriEmail or '')
                ws.cell(row=row, column=7, value=r.RandevuBaslik or '')
                ws.cell(row=row, column=8, value=r.Durum or '')
                ws.cell(row=row, column=9, value=f"{r.RandevuSuresi or 60} dk")
                ws.cell(row=row, column=10, value=r.defter.DefterAdi if r.defter else '')
            
            # Sütun genişliklerini ayarla
            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            # Excel dosyasını kaydet
            buffer = BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            
            filename = f"randevu_raporu_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
            return send_file(
                buffer,
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                as_attachment=True,
                download_name=filename
            )
        
        elif format_tip == 'pdf':
            # PDF Export
            buffer = BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18)
            
            # Türkçe font desteği
            try:
                pdfmetrics.registerFont(TTFont('DejaVuSans', 'C:/Windows/Fonts/dejavu-sans.ttf'))
                pdfmetrics.registerFont(TTFont('DejaVuSans-Bold', 'C:/Windows/Fonts/dejavu-sans-bold.ttf'))
                turkish_font = 'DejaVuSans'
                turkish_font_bold = 'DejaVuSans-Bold'
            except:
                try:
                    pdfmetrics.registerFont(TTFont('DejaVuSans', 'C:/Windows/Fonts/arial.ttf'))
                    pdfmetrics.registerFont(TTFont('DejaVuSans-Bold', 'C:/Windows/Fonts/arialbd.ttf'))
                    turkish_font = 'DejaVuSans'
                    turkish_font_bold = 'DejaVuSans-Bold'
                except:
                    turkish_font = 'Helvetica'
                    turkish_font_bold = 'Helvetica-Bold'
            
            # Stil tanımları
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=18, spaceAfter=30, alignment=TA_CENTER, fontName=turkish_font_bold)
            
            # Başlık
            title = Paragraph("Randevu Raporu", title_style)
            
            # Özet bilgiler
            toplam_randevu = len(randevular)
            tamamlanan = len([r for r in randevular if r.Durum == 'Tamamlandı'])
            iptal = len([r for r in randevular if r.Durum == 'İptal'])
            beklemede = len([r for r in randevular if r.Durum == 'Beklemede'])
            
            summary_data = [
                ['Toplam Randevu', str(toplam_randevu)],
                ['Tamamlanan', str(tamamlanan)],
                ['İptal', str(iptal)],
                ['Beklemede', str(beklemede)]
            ]
            
            summary_table = Table(summary_data)
            summary_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.lightgrey),
                ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, -1), turkish_font),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            # Tablo verileri
            table_data = [['Tarih', 'Saat', 'Müşteri', 'Telefon', 'Başlık', 'Durum', 'Süre']]
            
            for r in randevular[:50]:  # İlk 50 kayıt
                table_data.append([
                    r.RandevuTarihi.strftime('%d.%m.%Y'),
                    r.RandevuTarihi.strftime('%H:%M'),
                    f"{r.MusteriAdi} {r.MusteriSoyadi or ''}".strip(),
                    r.MusteriTelefon or '',
                    r.RandevuBaslik or '',
                    r.Durum or '',
                    f"{r.RandevuSuresi or 60} dk"
                ])
            
            # Tablo oluştur
            table = Table(table_data, repeatRows=1)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), turkish_font_bold),
                ('FONTSIZE', (0, 0), (-1, 0), 8),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('FONTNAME', (0, 1), (-1, -1), turkish_font),
                ('FONTSIZE', (0, 1), (-1, -1), 7),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            # PDF oluştur
            elements = [title, Spacer(1, 12), summary_table, Spacer(1, 12), table]
            doc.build(elements)
            buffer.seek(0)
            
            filename = f"randevu_raporu_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
            return send_file(
                buffer,
                mimetype='application/pdf',
                as_attachment=True,
                download_name=filename
            )

    return render_template('raporlar.html', 
                         baslangic=start_date.strftime('%Y-%m-%d'), 
                         bitis=end_date.strftime('%Y-%m-%d'),
                         kapasite=kapasite,
                         defter_id=defter_id,
                         defterler=defterler)



# Raporlar API - Ozet
@rapor_bp.route('/api/raporlar/ozet')
@login_required
def api_raporlar_ozet():
    baslangic = request.args.get('baslangic')
    bitis = request.args.get('bitis')
    defter_id = request.args.get('defter_id', '')
    try:
        start_date = datetime.strptime(baslangic, '%Y-%m-%d').date() if baslangic else (datetime.now().date() - timedelta(days=30))
        end_date = datetime.strptime(bitis, '%Y-%m-%d').date() if bitis else datetime.now().date()
    except ValueError:
        return jsonify({"success": False, "message": "Tarih formatı YYYY-MM-DD olmalı"}), 400

    # Tarihleri kapsayan datetime araligi
    start_dt = datetime(start_date.year, start_date.month, start_date.day, 0, 0)
    end_dt = datetime(end_date.year, end_date.month, end_date.day, 23, 59, 59)

    # Izin farkindaligi
    if session.get('is_admin', False):
        q = Randevu.query.filter(
            Randevu.FirmaID == session['firma_id'],
            Randevu.RandevuTarihi >= start_dt,
            Randevu.RandevuTarihi <= end_dt
        )
    else:
        q = db.session.query(Randevu).join(RandevuYetki).filter(
            RandevuYetki.KullaniciID == session['user_id'],
            RandevuYetki.GoruntulemeYetkisi == True,
            Randevu.FirmaID == session['firma_id'],
            Randevu.RandevuTarihi >= start_dt,
            Randevu.RandevuTarihi <= end_dt
        )
    
    # Defter filtresi ekle (Randevu.DefterID üzerinden)
    if defter_id:
        print(f"API Ozet - Defter filtresi uygulanıyor: {defter_id}")
        
        # Debug: Mevcut randevuları kontrol et
        # all_randevular = q.all()
        # print(f"DEBUG: Filtre öncesi toplam randevu: {len(all_randevular)}")
        
        # Debug: İlk randevunun DefterID'sini kontrol et
        # if all_randevular:
        #     first_randevu = all_randevular[0]
        #     print(f"DEBUG: İlk randevu DefterID: {first_randevu.DefterID}")
        
        # Debug: DefterID=5 olan randevuları kontrol et
        # defter_randevular = q.filter(Randevu.DefterID == defter_id).all()
        # print(f"DEBUG: DefterID={defter_id} olan randevu sayısı: {len(defter_randevular)}")
        
        q = q.filter(Randevu.DefterID == defter_id)
    else:
        print("API Ozet - Defter filtresi uygulanmıyor")

    items = q.all()
    print(f"API Ozet - Toplam {len(items)} randevu bulundu")

    toplam = len(items)
    durum_counter = Counter(r.Durum or 'Bilinmiyor' for r in items)
    ref_counter = Counter(r.RandevuBaslik or 'Yok' for r in items)

    saat_counter = defaultdict(int)
    for r in items:
        saat_counter[r.RandevuTarihi.hour] += 1
    saatler = list(range(0,24))
    saat_deger = [saat_counter.get(h, 0) for h in saatler]

    return jsonify({
        "success": True,
        "toplam": toplam,
        "durumlar": durum_counter,
        "referanslar": ref_counter,
        "saatler": {"labels": saatler, "values": saat_deger}
    })

# Yoğun Saatler Raporu API
@rapor_bp.route('/api/raporlar/yoğun-saatler')
@login_required
def api_raporlar_yogun_saatler():
    baslangic = request.args.get('baslangic')
    bitis = request.args.get('bitis')
    kapasite_param = request.args.get('kapasite', '')
    defter_id = request.args.get('defter_id', '')
    
    try:
        start_date = datetime.strptime(baslangic, '%Y-%m-%d').date() if baslangic else (datetime.now().date() - timedelta(days=30))
        end_date = datetime.strptime(bitis, '%Y-%m-%d').date() if bitis else datetime.now().date()
    except ValueError:
        return jsonify({"success": False, "message": "Tarih formatı YYYY-MM-DD olmalı"}), 400

    # Tarihleri kapsayan datetime araligi
    start_dt = datetime(start_date.year, start_date.month, start_date.day, 0, 0)
    end_dt = datetime(end_date.year, end_date.month, end_date.day, 23, 59, 59)

    # Eğer defter seçilmiş ve kapasite belirtilmemişse, defterden otomatik kapasite hesapla
    auto_capacity = None
    if defter_id and (not kapasite_param or kapasite_param == '0'):
        try:
            ayar = RandevuDefterAyar.query.filter_by(AyarID=defter_id, FirmaID=session['firma_id'], Aktif=True).first()
            if ayar:
                def parse_hhmm(s):
                    h, m = (s or '09:00').split(':')
                    return int(h), int(m)
                sh, sm = parse_hhmm(ayar.BaslangicSaati or '09:00')
                eh, em = parse_hhmm(ayar.BitisSaati or '18:00')
                total_minutes = max(0, (eh * 60 + em) - (sh * 60 + sm))
                slot_min = ayar.SlotDakika or 30
                auto_capacity = max(1, total_minutes // slot_min)
        except Exception:
            auto_capacity = None

    kapasite = int(kapasite_param) if (kapasite_param and kapasite_param != '0') else int(auto_capacity or 8)

    # Izin farkindaligi
    if session.get('is_admin', False):
        q = Randevu.query.filter(
            Randevu.FirmaID == session['firma_id'],
            Randevu.RandevuTarihi >= start_dt,
            Randevu.RandevuTarihi <= end_dt
        )
    else:
        q = db.session.query(Randevu).join(RandevuYetki).filter(
            RandevuYetki.KullaniciID == session['user_id'],
            RandevuYetki.GoruntulemeYetkisi == True,
            Randevu.FirmaID == session['firma_id'],
            Randevu.RandevuTarihi >= start_dt,
            Randevu.RandevuTarihi <= end_dt
        )
    
    # Defter filtresi ekle (Randevu.DefterID üzerinden)
    if defter_id:
        print(f"API Yogun Saatler - Defter filtresi uygulanıyor: {defter_id}")
        q = q.filter(Randevu.DefterID == defter_id)
    else:
        print("API Yogun Saatler - Defter filtresi uygulanmıyor")

    items = q.all()

    # Saatlik dağılım
    saat_counter = defaultdict(int)
    for r in items:
        saat_counter[r.RandevuTarihi.hour] += 1
    
    # 0-23 saat arası tüm saatler
    saatler = list(range(0, 24))
    saat_deger = [saat_counter.get(h, 0) for h in saatler]
    
    # En yoğun saat
    max_appointments = max(saat_deger) if saat_deger else 0
    peak_hour = saatler[saat_deger.index(max_appointments)] if max_appointments > 0 else 0
    
    # Ortalama kapasite kullanımı
    total_days = (end_date - start_date).days + 1
    total_possible_appointments = total_days * kapasite
    total_actual_appointments = len(items)
    avg_capacity_usage = (total_actual_appointments / total_possible_appointments * 100) if total_possible_appointments > 0 else 0
    
    # Verimlilik oranı (yoğun saatlerdeki verimlilik)
    peak_hours = [h for h in saatler if saat_counter.get(h, 0) >= max_appointments * 0.8]
    peak_hours_appointments = sum(saat_counter.get(h, 0) for h in peak_hours)
    efficiency_rate = (peak_hours_appointments / total_actual_appointments * 100) if total_actual_appointments > 0 else 0
    
    # Haftalık dağılım
    weekly_counter = defaultdict(int)
    for r in items:
        week_start = r.RandevuTarihi.date() - timedelta(days=r.RandevuTarihi.weekday())
        weekly_counter[week_start] += 1
    
    # Son 12 hafta
    weekly_labels = []
    weekly_values = []
    for i in range(12):
        week_start = end_date - timedelta(weeks=i)
        week_start = week_start - timedelta(days=week_start.weekday())
        weekly_labels.insert(0, week_start.strftime('%d/%m'))
        weekly_values.insert(0, weekly_counter.get(week_start, 0))
    
    # Saatlik verimlilik
    hourly_efficiency = []
    for h in saatler:
        hour_appointments = saat_counter.get(h, 0)
        hour_efficiency = (hour_appointments / max_appointments * 100) if max_appointments > 0 else 0
        hourly_efficiency.append(round(hour_efficiency, 1))

    return jsonify({
        "success": True,
        "peakHour": peak_hour,
        "avgCapacityUsage": round(avg_capacity_usage, 1),
        "efficiencyRate": round(efficiency_rate, 1),
        "hourlyData": {
            "labels": [f"{h:02d}" for h in saatler],
            "values": saat_deger
        },
        "weeklyData": {
            "labels": weekly_labels,
            "values": weekly_values
        },
        "hourlyEfficiency": hourly_efficiency
    })

@rapor_bp.route('/api/raporlar/heatmap')
@login_required
def api_raporlar_heatmap():
    baslangic = request.args.get('baslangic')
    bitis = request.args.get('bitis')
    defter_id = request.args.get('defter_id', '')
    
    try:
        start_date = datetime.strptime(baslangic, '%Y-%m-%d').date() if baslangic else (datetime.now().date() - timedelta(days=30))
        end_date = datetime.strptime(bitis, '%Y-%m-%d').date() if bitis else datetime.now().date()
    except ValueError:
        return jsonify({"success": False, "message": "Tarih formatı YYYY-MM-DD olmalı"}), 400

    # Tarihleri kapsayan datetime araligi
    start_dt = datetime(start_date.year, start_date.month, start_date.day, 0, 0)
    end_dt = datetime(end_date.year, end_date.month, end_date.day, 23, 59, 59)

    # Izin farkindaligi
    if session.get('is_admin', False):
        q = Randevu.query.filter(
            Randevu.FirmaID == session['firma_id'],
            Randevu.RandevuTarihi >= start_dt,
            Randevu.RandevuTarihi <= end_dt
        )
    else:
        q = db.session.query(Randevu).join(RandevuYetki).filter(
            RandevuYetki.KullaniciID == session['user_id'],
            RandevuYetki.GoruntulemeYetkisi == True,
            Randevu.FirmaID == session['firma_id'],
            Randevu.RandevuTarihi >= start_dt,
            Randevu.RandevuTarihi <= end_dt
        )
    
    # Defter filtresi ekle (Randevu.DefterID üzerinden)
    if defter_id:
        print(f"API Heatmap - Defter filtresi uygulanıyor: {defter_id}")
        q = q.filter(Randevu.DefterID == defter_id)
    else:
        print("API Heatmap - Defter filtresi uygulanmıyor")

    items = q.all()

    # Günlük ve saatlik dağılım
    heatmap_data = {
        'Monday': {},
        'Tuesday': {},
        'Wednesday': {},
        'Thursday': {},
        'Friday': {},
        'Saturday': {},
        'Sunday': {}
    }
    
    # Gün isimleri
    day_names = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
    
    for r in items:
        # Gün adını al (0=Monday, 6=Sunday)
        day_name = day_names[r.RandevuTarihi.weekday()]
        hour = r.RandevuTarihi.hour
        
        # Sadece çalışma saatleri (9-18)
        if 9 <= hour <= 18:
            if hour not in heatmap_data[day_name]:
                heatmap_data[day_name][hour] = 0
            heatmap_data[day_name][hour] += 1
    
    return jsonify({
        "success": True,
        "heatmapData": heatmap_data
    })

# Personel Performans Raporu API
@rapor_bp.route('/api/raporlar/personel-performans')
@login_required
def api_raporlar_personel_performans():
    baslangic = request.args.get('baslangic')
    bitis = request.args.get('bitis')
    defter_id = request.args.get('defter_id', '')

    try:
        start_date = datetime.strptime(baslangic, '%Y-%m-%d').date() if baslangic else (datetime.now().date() - timedelta(days=30))
        end_date = datetime.strptime(bitis, '%Y-%m-%d').date() if bitis else datetime.now().date()
    except ValueError:
        return jsonify({"success": False, "message": "Tarih formatı YYYY-MM-DD olmalı"}), 400

    start_dt = datetime(start_date.year, start_date.month, start_date.day, 0, 0)
    end_dt = datetime(end_date.year, end_date.month, end_date.day, 23, 59, 59)

    # İzin farkındalığı
    if session.get('is_admin', False):
        q = Randevu.query.filter(
            Randevu.FirmaID == session['firma_id'],
            Randevu.RandevuTarihi >= start_dt,
            Randevu.RandevuTarihi <= end_dt
        )
    else:
        q = db.session.query(Randevu).join(RandevuYetki).filter(
            RandevuYetki.KullaniciID == session['user_id'],
            RandevuYetki.GoruntulemeYetkisi == True,
            Randevu.FirmaID == session['firma_id'],
            Randevu.RandevuTarihi >= start_dt,
            Randevu.RandevuTarihi <= end_dt
        )

    if defter_id:
        q = q.filter(Randevu.DefterID == defter_id)

    items = q.all()

    # Personel bazında grupla
    performance = {}
    for r in items:
        user_id = r.OlusturanKullaniciID
        if user_id not in performance:
            performance[user_id] = {
                'kullaniciId': user_id,
                'adSoyad': f"{r.olusturan_kullanici.Ad} {r.olusturan_kullanici.Soyad}" if r.olusturan_kullanici else 'Bilinmiyor',
                'toplam': 0,
                'durumlar': defaultdict(int)
            }
        performance[user_id]['toplam'] += 1
        durum = r.Durum or 'Bilinmiyor'
        performance[user_id]['durumlar'][durum] += 1

    # Sonuçları listeye çevir ve toplam sayıya göre sırala
    rows = []
    for _, info in performance.items():
        rows.append({
            'kullaniciId': info['kullaniciId'],
            'adSoyad': info['adSoyad'],
            'toplam': info['toplam'],
            'beklemede': info['durumlar'].get('Beklemede', 0),
            'tamamlandi': info['durumlar'].get('Tamamlandı', 0),
            'iptal': info['durumlar'].get('İptal', 0) + info['durumlar'].get('Iptal', 0)
        })
    rows.sort(key=lambda x: x['toplam'], reverse=True)

    labels = [r['adSoyad'] for r in rows]
    counts = [r['toplam'] for r in rows]

    return jsonify({
        'success': True,
        'labels': labels,
        'counts': counts,
        'rows': rows
    })
